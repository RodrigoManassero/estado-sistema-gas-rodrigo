#!/usr/bin/env python3
"""Fetch CAMMESA Parte Post Operativo (PPO) — daily fuel mix + generation.

The PPO is published daily via the public API at api.cammesa.com/pub-svc/public.
Each file is an OLE2 .xls "protected" with the Microsoft default read-only
password VelvetSweatshop (no real authentication — Excel opens these files
without prompting). We decrypt, aggregate the "Consumo Combustibles" sheet
across every plant, and write a daily time series into cammesa_ppo.json.

Units:
- Gas natural: dam³ (1 dam³ = 1000 m³, so divide by 1000 to get MMm³/d)
- Carbón mineral: Tn
- Fuel Oil: Tn
- Gas Oil: m³
- Generación neta: MWh

Usage:
    python scripts/fetch_cammesa_ppo.py --days 30
    python scripts/fetch_cammesa_ppo.py --days 365 --force
"""

import argparse
import io
import json
import os
import sys
from datetime import date, datetime, timedelta

import requests

try:
    import msoffcrypto
    import xlrd
except ImportError as e:
    print(f"Install dependencies first: pip install msoffcrypto-tool xlrd  ({e})", file=sys.stderr)
    sys.exit(1)

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _meta import write_json, write_csv, json_to_csv_path  # noqa: E402

OUT_DIR = os.path.join(os.path.dirname(__file__), '..', 'public', 'data')
CAMMESA_JSON = os.path.join(OUT_DIR, 'cammesa_ppo.json')

API_BASE = 'https://api.cammesa.com/pub-svc/public/'
NEMO = 'INFORME_OPERATIVO'
DEFAULT_PW = 'VelvetSweatshop'  # Excel's default read-only password
HDRS = {'User-Agent': 'Mozilla/5.0 Chrome/130 estado-red-gas'}


def load_existing():
    if not os.path.exists(CAMMESA_JSON):
        return {}
    with open(CAMMESA_JSON, encoding='utf-8') as f:
        raw = json.load(f)
    data = raw.get('data', raw) if isinstance(raw, dict) else raw
    return {r['fecha']: r for r in (data or []) if r.get('fecha')}


def list_documents(date_from: date, date_to: date, session: requests.Session):
    """Query findDocumentosByNemoRango. Returns list of doc dicts with attachments."""
    params = {
        'nemo': NEMO,
        'fechadesde': f'{date_from.isoformat()}T00:00:00',
        'fechahasta': f'{date_to.isoformat()}T23:59:59',
    }
    r = session.get(API_BASE + 'findDocumentosByNemoRango', params=params, headers=HDRS, timeout=60)
    r.raise_for_status()
    print("SIZE:", len(r.content))
    print("FIRST BYTES:", r.content[:16])
    return r.json()


def download_and_decrypt(doc_id: str, attachment_id: str, session: requests.Session) -> bytes:
    r = session.get(API_BASE + 'findAttachmentByNemoId', params={
        'attachmentId': attachment_id,
        'docId': doc_id,
        'nemo': NEMO,
    }, headers=HDRS, timeout=120)
    r.raise_for_status()
    print("CONTENT TYPE:", r.headers.get("content-type"))
    print("SIZE:", len(r.content))
    print("FIRST BYTES:", r.content[:16])
    
    # XLSX moderno (OpenXML)
    if r.content[:4] == b'PK\x03\x04':
        return r.content

    # XLS clásico protegido
    if r.content[:4] == b'\xd0\xcf\x11\xe0':
        office = msoffcrypto.OfficeFile(io.BytesIO(r.content))
        try:
            office.load_key(password=DEFAULT_PW)
        except Exception as e:
            raise ValueError(f'decrypt failed: {e}')

        out = io.BytesIO()
        office.decrypt(out)
        out.seek(0)
        return out.read()

    raise ValueError(
        f'unsupported file type (first bytes {r.content[:8]!r})'
    )


def aggregate_fuel_sheet(xls_bytes: bytes):
    """Sum every plant row in Consumo Combustibles, return totals + metadata."""

    # XLSX moderno
    if xls_bytes[:4] == b'PK\x03\x04':
        from openpyxl import load_workbook

        wb = load_workbook(
            io.BytesIO(xls_bytes),
            data_only=True
        )

        if "Consumo Combustibles" not in wb.sheetnames:
            return None

        ws = wb["Consumo Combustibles"]

        totals = {
            'gas_dam3': 0.0,
            'carbon_tn': 0.0,
            'fueloil_tn': 0.0,
            'gasoil_m3': 0.0,
            'gen_gas_mwh': 0.0,
            'gen_carbon_mwh': 0.0,
            'gen_fueloil_mwh': 0.0,
            'gen_gasoil_mwh': 0.0,
        }
        plants_counted = 0

        # Sumar desde la fila 5 hasta la 299 (o max_row)
        max_r = min(299, ws.max_row)
        for r in range(5, max_r + 1):
            code = str(ws.cell(row=r, column=2).value or '').strip()
            if not code or code.lower().startswith('total'):
                continue
            def num(col_idx):
                v = ws.cell(row=r, column=col_idx).value
                try:
                    return float(v or 0)
                except (TypeError, ValueError):
                    return 0.0

            totals['gas_dam3'] += num(3)          # Columna C
            totals['gen_gas_mwh'] += num(4)       # Columna D
            totals['carbon_tn'] += num(6)         # Columna F
            totals['gen_carbon_mwh'] += num(7)    # Columna G
            totals['fueloil_tn'] += num(9)        # Columna I
            totals['gen_fueloil_mwh'] += num(10)  # Columna J
            totals['gasoil_m3'] += num(12)        # Columna L
            totals['gen_gasoil_mwh'] += num(13)   # Columna M
            plants_counted += 1

        # Intentar extraer la fecha de la hoja (ej. fila 2)
        fecha_iso = None
        for r in range(1, 4):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, str) and 'Día' in v:
                    try:
                        fecha_iso = datetime.strptime(v.split(':', 1)[1].strip(), '%d/%m/%Y').date().isoformat()
                    except ValueError:
                        pass
                    break
            if fecha_iso:
                break

        return {
            'fecha': fecha_iso,
            'gas_mmm3': round(totals['gas_dam3'] / 1000, 3),  # dam³ -> MMm³
            'carbon_tn': round(totals['carbon_tn'], 1),
            'fueloil_tn': round(totals['fueloil_tn'], 1),
            'gasoil_m3': round(totals['gasoil_m3'], 1),
            'gen_gas_mwh': round(totals['gen_gas_mwh'], 1),
            'gen_carbon_mwh': round(totals['gen_carbon_mwh'], 1),
            'gen_fueloil_mwh': round(totals['gen_fueloil_mwh'], 1),
            'gen_gasoil_mwh': round(totals['gen_gasoil_mwh'], 1),
            'plants_counted': plants_counted,
        }

    # XLS clásico
    wb = xlrd.open_workbook(
        file_contents=xls_bytes,
        on_demand=True
    )

    if 'Consumo Combustibles' not in wb.sheet_names():
        return None

    s = wb.sheet_by_name('Consumo Combustibles')
    fecha_iso = None
    for r in range(3):
        for c in range(s.ncols):
            v = s.cell_value(r, c)
            if isinstance(v, str) and 'Día' in v:
                try:
                    fecha_iso = datetime.strptime(v.split(':', 1)[1].strip(), '%d/%m/%Y').date().isoformat()
                except ValueError:
                    pass
                break
        if fecha_iso:
            break

    totals = {
        'gas_dam3': 0.0,
        'carbon_tn': 0.0,
        'fueloil_tn': 0.0,
        'gasoil_m3': 0.0,
        'gen_gas_mwh': 0.0,
        'gen_carbon_mwh': 0.0,
        'gen_fueloil_mwh': 0.0,
        'gen_gasoil_mwh': 0.0,
    }
    plants_counted = 0
    max_r = min(299, s.nrows)
    for r in range(4, max_r):
        code = str(s.cell_value(r, 1) or '').strip()
        if not code or code.lower().startswith('total'):
            continue
        def num(col):
            v = s.cell_value(r, col)
            try:
                return float(v or 0)
            except (TypeError, ValueError):
                return 0.0
        totals['gas_dam3'] += num(2)
        totals['gen_gas_mwh'] += num(3)
        totals['carbon_tn'] += num(5)
        totals['gen_carbon_mwh'] += num(6)
        totals['fueloil_tn'] += num(8)
        totals['gen_fueloil_mwh'] += num(9)
        totals['gasoil_m3'] += num(11)
        totals['gen_gasoil_mwh'] += num(12)
        plants_counted += 1

    return {
        'fecha': fecha_iso,
        'gas_mmm3': round(totals['gas_dam3'] / 1000, 3),  # dam³ -> MMm³
        'carbon_tn': round(totals['carbon_tn'], 1),
        'fueloil_tn': round(totals['fueloil_tn'], 1),
        'gasoil_m3': round(totals['gasoil_m3'], 1),
        'gen_gas_mwh': round(totals['gen_gas_mwh'], 1),
        'gen_carbon_mwh': round(totals['gen_carbon_mwh'], 1),
        'gen_fueloil_mwh': round(totals['gen_fueloil_mwh'], 1),
        'gen_gasoil_mwh': round(totals['gen_gasoil_mwh'], 1),
        'plants_counted': plants_counted,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--days', type=int, default=10, help='days back from today')
    parser.add_argument('--force', action='store_true', help='re-fetch dates we already have')
    parser.add_argument('--sleep-ms', type=int, default=150)
    args = parser.parse_args()

    os.makedirs(OUT_DIR, exist_ok=True)
    existing = load_existing()
    print(f"Starting with {len(existing)} existing rows in cammesa_ppo.json")

    today = date.today()
    date_from = today - timedelta(days=args.days)

    session = requests.Session()
    print(f"Listing PPO documents {date_from} to {today}...")
    try:
        docs = list_documents(date_from, today, session)
    except Exception as e:
        print(f"ERROR listing: {e}", file=sys.stderr)
        return 1

    print(f"  {len(docs)} documents found")

    added = 0
    failures = []
    for doc in docs:
        try:
            fecha_iso = datetime.strptime(doc['fecha'], '%d/%m/%Y').date().isoformat()
        except (KeyError, ValueError):
            failures.append(f"bad fecha in doc {doc.get('id')}")
            continue

        if fecha_iso in existing and not args.force:
            continue

        attachments = doc.get('adjuntos') or []
        if not attachments:
            failures.append(f"{fecha_iso}: no attachments")
            continue
        att = attachments[0]
        attachment_id = att.get('id') or att.get('nombre')
        if not attachment_id:
            failures.append(f"{fecha_iso}: missing attachment id")
            continue

        try:
            xls = download_and_decrypt(doc['id'], attachment_id, session)
            aggregated = aggregate_fuel_sheet(xls)
        except Exception as e:
            failures.append(f"{fecha_iso}: {e}")
            print(f"  FAIL {fecha_iso}: {e}", file=sys.stderr)
            continue

        if not aggregated:
            failures.append(f"{fecha_iso}: no Consumo Combustibles sheet")
            continue
            
        aggregated['fecha'] = aggregated.get('fecha') or fecha_iso
        aggregated['source'] = attachment_id
        existing[aggregated['fecha']] = aggregated
        added += 1
        if added % 5 == 0:
            print(f"  +{added} gas={aggregated['gas_mmm3']} MMm³")
        if args.sleep_ms:
            import time as _t
            _t.sleep(args.sleep_ms / 1000)

    rows = sorted(existing.values(), key=lambda r: r.get('fecha') or '')
    latest = rows[-1]['fecha'] if rows else None
    write_json(
        CAMMESA_JSON,
        rows,
        source='CAMMESA PPO — Consumo Combustibles (aggregated per day)',
        source_date=latest,
        failures=failures[-30:],
    )
    ppo_cols = [
        'fecha', 'source', 'gas_mmm3', 'carbon_tn', 'fueloil_tn', 'gasoil_m3',
        'gen_gas_mwh', 'gen_carbon_mwh', 'gen_fueloil_mwh', 'gen_gasoil_mwh',
        'plants_counted',
    ]
    write_csv(json_to_csv_path(CAMMESA_JSON), rows, fieldnames=ppo_cols)
    print(f"\ncammesa_ppo.json: {len(rows)} rows ({added} new, {len(failures)} failed)")
    return 0


if __name__ == '__main__':
    sys.exit(main())
