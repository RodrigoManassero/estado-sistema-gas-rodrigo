#!/usr/bin/env python3
"""Parse CCO Word minutas (.docx) from raw/ and generate clean public/data/mantenimientos.json."""

import json
import os
import re
import sys
from datetime import datetime
import docx

OUT_DIR = os.path.join(os.path.dirname(__file__), '..', 'public', 'data')
RAW_DIR = os.path.join(os.path.dirname(__file__), '..', 'raw')
OUT_JSON = os.path.join(OUT_DIR, 'mantenimientos.json')


def clean_text(text: str) -> str:
    return ' '.join(text.split())


def extract_impact_text(text: str) -> str:
    match = re.search(r'(?:afectación|impacto|impacto potencial|impacto estimado)[:\s]*(.*)', text, re.IGNORECASE)
    if match:
        val_str = clean_text(match.group(1))
        val_lower = val_str.lower()
        if any(kw in val_lower for kw in ['sin impacto', 'no', 'sin restricción', 'sin afectación', 'no hay']):
            return "0 MM³/día"
        return val_str
        
    text_lower = text.lower()
    if any(kw in text_lower for kw in ['sin impacto', 'sin restricción', 'sin afectación']):
        return "0 MM³/día"
        
    return "0 MM³/día"


def extract_month(text: str) -> str:
    text_lower = text.lower()
    if 'octubre' in text_lower or '10/' in text_lower or '28/10' in text_lower:
        return "2026-10"
    if 'noviembre' in text_lower or '11/' in text_lower:
        return "2026-11"
    if 'diciembre' in text_lower or '12/' in text_lower:
        return "2026-12"
    
    date_match = re.search(r'(\d{1,2})/(\d{1,2})', text)
    if date_match:
        m = int(date_match.group(2))
        if 1 <= m <= 12:
            return f"2026-{m:02d}"
            
    return "2026-10"


def parse_fecha_y_duracion(text: str) -> tuple[str, str]:
    text_clean = clean_text(text)
    
    date_match = re.search(r'(\d{1,2}/\d{1,2})', text_clean)
    if date_match and not "2026" in text_clean[:5]:
        fecha = date_match.group(1)
    else:
        fecha = "A programar"
        
    dur_match = re.search(r'(\d+\s*(?:d[ií]as?|hs?|horas?))', text_clean, re.IGNORECASE)
    duracion = dur_match.group(1) if dur_match else "A confirmar"
    
    return fecha, duracion


def simplify_obra_tarea(text: str) -> str:
    text_lower = text.lower()
    if 'limpieza' in text_lower and not ('instrumentado' in text_lower or 'tfi' in text_lower or 'mfl' in text_lower):
        return "Scraper Limpieza"
    elif any(kw in text_lower for kw in ['instrumentado', 'tfi', 'mfl', 'clp']):
        return "Scraper instrumentado"
    elif any(kw in text_lower for kw in ['corridas', 'corrida de scrapper', 'corrida de scraper', 'corridas de scrapper']):
        return "Scraper limpieza + instrumentado"
    return text


def parse_cco_docx(file_path: str) -> list[dict]:
    doc = docx.Document(file_path)
    records = []
    
    current_cuenca = "Cuenca Neuquina"
    current_empresa = "TGS"
    current_instalacion = "General"
    
    ignore_phrases = [
        "comité de coordinación", "minuta reunión", "lugar:", "participantes:", 
        "se llevó a cabo", "reunión mensual", "no tiene mantenimientos", 
        "sin novedades", "no participó", "no realizó comentarios", 
        "se continuará con las reuniones", "los mantenimientos previstos fueron realizados", 
        "se finalizaron todas las tareas", "fallas (revisión y reparación)", 
        "campaña verificación de defectos", "obras en gasoducto", "obras en plantas",
        "trabajos en plantas"
    ]

    current_item = None
    in_operational_data = False

    for p in doc.paragraphs:
        text = clean_text(p.text)
        if not text or len(text) < 3:
            continue
            
        text_lower = text.lower()
        
        # Ignorar encabezados y títulos de categoría
        if any(phrase in text_lower for phrase in ignore_phrases):
            continue
            
        # Detectar Secciones / Cuencas
        if "Cuenca" in text or "Tramos Finales" in text or "Gasoducto" in text:
            if len(text) < 40:
                # Si teníamos un item previo en curso, guardarlo
                if current_item and current_item.get("obra_tarea"):
                    records.append(current_item)
                    current_item = None
                    
                current_cuenca = text
                in_operational_data = True
                continue
                
        if not in_operational_data and not any(e in text.upper() for e in ["TGS", "TGN"]):
            continue

        # Cambiar de Empresa (Cierra el item actual para no contaminar valores)
        if text.upper() in ["TGS", "TGN"]:
            if current_item and current_item.get("obra_tarea"):
                records.append(current_item)
                current_item = None
            current_empresa = text.upper()
            continue

        # Detectar nueva instalación (Tramo o Planta)
        if text.startswith("Tramo ") or text.startswith("Planta ") or text.startswith("Sección "):
            if current_item and current_item.get("obra_tarea"):
                records.append(current_item)
                current_item = None
                
            # Limpiar dos puntos al final si los tiene (ej: "Planta Mercedes:")
            current_instalacion = text.rstrip(":")
            continue

        # Iniciar bloque independiente por Obra o Tarea
        if text.startswith("Obra:") or text.startswith("Tarea:"):
            if current_item and current_item.get("obra_tarea"):
                records.append(current_item)
            
            raw_obra = text.split(":", 1)[1].strip()
            current_item = {
                "cuenca_zona": current_cuenca,
                "empresa": current_empresa,
                "instalacion": current_instalacion,
                "obra_tarea": simplify_obra_tarea(raw_obra),
                "fecha": "A programar",
                "duracion": "A confirmar",
                "mes_evento": "2026-10",
                "impacto_mmm3d": "0 MM³/día"
            }
            continue

        if not current_item:
            current_item = {
                "cuenca_zona": current_cuenca,
                "empresa": current_empresa,
                "instalacion": current_instalacion,
                "obra_tarea": "",
                "fecha": "A programar",
                "duracion": "A confirmar",
                "mes_evento": "2026-10",
                "impacto_mmm3d": "0 MM³/día"
            }

        if any(kw in text_lower for kw in ['cambio de', 'reparación', 'prueba', 'corrida', 'paro', 'mantenimiento', 'verificación', 'adecuación', 'tie-in', 'recobertura', 'pasajes de scraper']) and not current_item["obra_tarea"]:
            current_item["obra_tarea"] = simplify_obra_tarea(text)

        # Capturar Fecha / Periodo
        if text.startswith("Fecha:") or text.startswith("Periodo:"):
            fecha_val = text.split(":", 1)[1].strip()
            f_ext, d_ext = parse_fecha_y_duracion(fecha_val)
            current_item["fecha"] = f_ext
            current_item["duracion"] = d_ext
            current_item["mes_evento"] = extract_month(fecha_val)

        # Capturar Afectación / Impacto
        if text.startswith("Afectación:") or text.startswith("Impacto:") or text.startswith("Impacto potencial:") or text.startswith("Impacto estimado:"):
            current_item["impacto_mmm3d"] = extract_impact_text(text)

    # Cierre del último item del bucle de párrafos
    if current_item and current_item.get("obra_tarea"):
        records.append(current_item)

    # ------------------------------------------------------------------------
    # Procesamiento de Tablas (GPM y acometidas) + Deduplicación Inteligente
    # ------------------------------------------------------------------------
    existing_plants = {r["instalacion"].lower() for r in records}

    for table in doc.tables:
        for row in table.rows:
            cells = [clean_text(cell.text) for cell in row.cells]
            if len(cells) >= 3:
                inst_raw, fecha_raw, imp_raw = cells[0], cells[1], cells[2]
                
                # Omitir cabecera de la tabla
                if inst_raw and fecha_raw and not any(h in inst_raw.lower() for h in ['planta', 'empresa', 'apellido', 'instalación']):
                    
                    # Normalizar nombre de planta (ej. "Casa de Piedra" -> "Planta Casa de Piedra")
                    inst_name = inst_raw if inst_raw.startswith("Planta ") else f"Planta {inst_raw}"
                    
                    # Deduplicación: Si la planta ya fue parseada detalladamente en los párrafos, la salteamos
                    if any(inst_raw.lower() in existing_p or existing_p in inst_raw.lower() for existing_p in existing_plants):
                        continue
                        
                    f_ext, d_ext = parse_fecha_y_duracion(fecha_raw)
                    records.append({
                        "cuenca_zona": "Gasoducto GPM",
                        "empresa": "TGS",
                        "instalacion": inst_name,
                        "obra_tarea": "Acometidas para futuras plantas compresoras",
                        "fecha": f_ext,
                        "duracion": d_ext,
                        "mes_evento": extract_month(fecha_raw),
                        "impacto_mmm3d": extract_impact_text(f"Impacto: {imp_raw}")
                    })

    return records


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    
    if not os.path.exists(RAW_DIR):
        print(f"ERROR: La carpeta {RAW_DIR} no existe.", file=sys.stderr)
        sys.exit(1)
        
    doc_files = [os.path.join(RAW_DIR, f) for f in os.listdir(RAW_DIR) if f.startswith("Minuta") and f.endswith(".docx")]
    if not doc_files:
        print("ERROR: No se encontró ningún archivo 'Minuta...docx' en raw/", file=sys.stderr)
        sys.exit(1)
        
    latest_file = max(doc_files, key=os.path.getmtime)
    print(f"Procesando minuta: {latest_file}")
    
    data = parse_cco_docx(latest_file)
    
    output = {
        "generated_at": datetime.now().isoformat(),
        "source": os.path.basename(latest_file),
        "data": data
    }
    
    with open(OUT_JSON, 'w', encoding='utf-8') as f:
        json.dump(output, f, indent=2, ensure_ascii=False)
        
    print(f"ÉXITO: Generado {OUT_JSON} con {len(data)} registros estructurados.")


if __name__ == '__main__':
    main()
